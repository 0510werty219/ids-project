# -*- coding: utf-8 -*-
"""DOCX 版式工具：样式、表格、图表、富文本。"""
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x52, 0x60, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD_FILL = "2F5F7F"
ALT_FILL = "F2F6F9"
BORDER = "D9D9D9"
BODY_FONT = "Times New Roman"
BODY_EA = "宋体"
UI_FONT = "Microsoft YaHei"
UI_EA = "微软雅黑"


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


def num(x):
    return f"{int(x):,}"


def set_run_font(run, name=BODY_FONT, east=BODY_EA, size=10.5, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def cell_margins(table, top=60, start=100, bottom=60, end=100):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


class Builder:
    def __init__(self, doc):
        self.doc = doc

    def add_rich(self, paragraph, text, size=10.5, name=BODY_FONT, east=BODY_EA, color=INK):
        for i, seg in enumerate(text.split("**")):
            if not seg:
                continue
            run = paragraph.add_run(seg)
            set_run_font(run, name=name, east=east, size=size, bold=(i % 2 == 1), color=color)
        return paragraph

    def setup_styles(self):
        doc = self.doc
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)

        normal = doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_EA)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal.paragraph_format.line_spacing = 1.3
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.space_before = Pt(0)

        title = doc.styles["Title"]
        title.font.name = UI_FONT
        title.font.size = Pt(20)
        title.font.bold = True
        title.font.color.rgb = INK
        title.element.rPr.rFonts.set(qn("w:eastAsia"), UI_EA)
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(4)
        title.paragraph_format.line_spacing = 1.15
        # Word 的 Title 样式自带下边框线，按版式要求去掉
        tpr = title.element.get_or_add_pPr()
        pbd = tpr.find(qn("w:pBdr"))
        if pbd is not None:
            tpr.remove(pbd)

        for style_name, size, before, after in (("Heading 1", 15, 18, 8),
                                                ("Heading 2", 12, 14, 5),
                                                ("Heading 3", 10.5, 10, 4)):
            st = doc.styles[style_name]
            st.font.name = UI_FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = INK
            st.element.rPr.rFonts.set(qn("w:eastAsia"), UI_EA)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.line_spacing = 1.2

    def body(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10.5, space_after=7):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        self.add_rich(p, text, size=size)
        return p

    def bullet(self, text, size=10.5):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.left_indent = Cm(0.75)
        self.add_rich(p, text, size=size)
        return p

    def numbered(self, idx, text, size=10.5):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        self.add_rich(p, f"{idx}. {text}", size=size)
        return p

    def table_caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        self.add_rich(p, text, size=9.5, name=UI_FONT, east=UI_EA, color=INK)
        return p

    def figure(self, path, caption, width_cm=15.5):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        pic = p.add_run().add_picture(path, width=Cm(width_cm))
        # 可访问性：为图片补充替代文字
        doc_pr = pic._inline.docPr
        doc_pr.set("descr", caption)
        doc_pr.set("title", caption)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(12)
        self.add_rich(cap, caption, size=9, name=UI_FONT, east=UI_EA, color=GRAY)

    def make_table(self, headers, rows, widths, aligns=None, font_size=9.5,
                   head_font_size=9.5):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table_borders(table)
        cell_margins(table)
        # 网格列宽与单元格宽度都要设置，否则 Word 按等宽列渲染
        for i, w in enumerate(widths):
            table.columns[i].width = Cm(w)
        aligns = aligns or ["left"] * len(headers)
        align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT}
        hdr = table.rows[0]
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            shade(cell, HEAD_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(h)
            set_run_font(run, name=UI_FONT, east=UI_EA, size=head_font_size, bold=True,
                         color=WHITE)
        repeat_header(hdr)
        for r_idx, row in enumerate(rows):
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cell = cells[i]
                cell.width = Cm(widths[i])
                if r_idx % 2 == 1:
                    shade(cell, ALT_FILL)
                p = cell.paragraphs[0]
                p.alignment = align_map[aligns[i]]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                self.add_rich(p, str(val), size=font_size)
        return table

    def spacer(self, pts=4):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(pts)
        p.paragraph_format.space_before = Pt(0)
        return p
